
import os
import json
from pathlib import Path

from dotenv import load_dotenv
from google import genai
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

API_KEY = os.getenv("GEMINI_API_KEY")
if not API_KEY:
    raise ValueError("GEMINI_API_KEY missing from .env")

client = genai.Client(api_key=API_KEY)
MODEL_NAME = "gemini-3.5-flash-lite"


def extract_text(file_path: str) -> str:
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    if extension == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if extension == ".docx":
        document = Document(str(path))
        return "\n".join(p.text for p in document.paragraphs)

    if extension == ".pptx":
        presentation = Presentation(str(path))
        slides_text = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slides_text.append(shape.text)
        return "\n".join(slides_text)

    raise ValueError("Supported formats: PDF, DOCX, PPTX, TXT")


def generate_quiz(material_text: str, question_count: int = 5) -> list:
    if not material_text.strip():
        raise ValueError(
            "No readable text found in this file. "
            "Scanned PDFs may need OCR."
        )

    # Keep the prompt within a practical size for this first version.
    material_text = material_text[:25000]

    prompt = f"""
Create exactly {question_count} multiple-choice questions
based ONLY on the study material below.

Return a JSON array. Each item must contain:
- question: a string
- options: exactly four strings
- correct_answer: the exact correct option string
- explanation: a short explanation

Make the questions clear and educational.
Do not include information that is not supported by the material.

STUDY MATERIAL:
{material_text}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config={
            "response_mime_type": "application/json"
        }
    )

    questions = json.loads(response.text)

    if not isinstance(questions, list) or not questions:
        raise ValueError("AI did not return a valid question list.")

    for question in questions:
        if not isinstance(question.get("question"), str):
            raise ValueError("Invalid question returned by AI.")
        if not isinstance(question.get("options"), list):
            raise ValueError("Invalid options returned by AI.")
        if len(question["options"]) != 4:
            raise ValueError("Each question must have four options.")
        if question.get("correct_answer") not in question["options"]:
            raise ValueError("Correct answer is not one of the options.")

    return questions
